import unittest
from pathlib import Path
import sys

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from modules.speech.services.insight_lead_service import InsightLeadService
from modules.verint.services.verint_api_client import VerintAPIClient


class TestSpeechAndVerintModules(unittest.TestCase):

    def test_insight_query_builder(self):
        """Verifica que el SQL dinámico de Insight incluya los IDs y la estructura de tablas temporales."""
        service = InsightLeadService(username="dummy", password="dummy")
        test_ids = ["CONID-123", "CONID-456"]
        sql = service._build_query_sql(test_ids, min_date="2026-07-01")

        self.assertIn("'CONID-123'", sql)
        self.assertIn("'CONID-456'", sql)
        self.assertIn("INTO #BASE", sql)
        self.assertIn("INTO #HISTORICO", sql)
        self.assertIn("2026-07-01", sql)
        self.assertIn("TIPO_LEAD", sql)

    def test_verint_format_dialogue_success(self):
        """Verifica el formateo correcto de respuestas JSON de Verint a texto plano con minutaje."""
        mock_response = {
            "GetInteractionTranscriptionResult": {
                "Success": True,
                "Data": {
                    "WordsSequences": [
                        {
                            "SpeakerName": "Agent",
                            "StartTime": 0,
                            "Words": [{"WordText": "Hola"}, {"WordText": "buenas"}, {"WordText": "tardes"}]
                        },
                        {
                            "SpeakerName": "Customer",
                            "StartTime": 65000,
                            "Words": [{"WordText": "Hola"}, {"WordText": "gracias"}]
                        }
                    ]
                }
            }
        }

        dialogue = VerintAPIClient.format_dialogue(mock_response)
        lines = dialogue.splitlines()

        self.assertEqual(len(lines), 2)
        self.assertEqual(lines[0], "[00:00] Asesor: Hola buenas tardes")
        self.assertEqual(lines[1], "[01:05] Cliente: Hola gracias")

    def test_verint_format_dialogue_empty_or_null(self):
        """Verifica que respuestas vacías o inválidas no rompan la ejecución."""
        self.assertEqual(VerintAPIClient.format_dialogue(None), "")
        self.assertEqual(VerintAPIClient.format_dialogue({}), "")
    def test_export_verint_docx_generation(self):
        """Verifica la generación correcta del documento Word (.docx) a partir de eventos de Verint."""
        import tempfile
        from modules.verint.tools.export_verint_docx_by_ids import (
            extract_call_dialogue_events,
            create_interaction_docx
        )
        import docx

        mock_response = {
            "GetInteractionTranscriptionResult": {
                "Success": True,
                "Data": {
                    "WordsSequences": [
                        {
                            "SpeakerName": "Agent",
                            "StartTime": 0,
                            "Words": [{"WordText": "Buenos"}, {"WordText": "días"}]
                        },
                        {
                            "SpeakerName": "Customer",
                            "StartTime": 12000,
                            "Words": [{"WordText": "Hola"}]
                        }
                    ]
                }
            }
        }
        events = extract_call_dialogue_events(mock_response, "Juan Perez")
        self.assertEqual(len(events), 2)
        self.assertEqual(events[0]["tipo"], "Interno")
        self.assertEqual(events[1]["tipo"], "Externo")

        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = Path(tmpdir) / "test_interaction.docx"
            meta = {"Agent": "Juan Perez", "LocalStartTime": "2026-09-10 10:00:00", "SID": "12345"}
            created_path = create_interaction_docx("99998888", meta, events, out_file)
            self.assertTrue(created_path.exists())

            # Validar que python-docx puede leer el archivo y sus tablas
            doc = docx.Document(str(created_path))
            self.assertGreaterEqual(len(doc.tables), 2)
            self.assertIn("99998888", doc.paragraphs[0].text)


if __name__ == "__main__":
    unittest.main()
