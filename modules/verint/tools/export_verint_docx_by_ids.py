import os
import sys
import argparse
import logging
from pathlib import Path
from typing import List, Dict, Any, Optional
from dotenv import load_dotenv

BASE_DIR = Path(__file__).resolve().parents[3]
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

import docx
from docx.shared import Pt, RGBColor
from modules.verint.services.verint_api_client import VerintAPIClient

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S"
)
logger = logging.getLogger("VerintDocxExporter")


def extract_call_dialogue_events(res_data: Dict[str, Any], agent_name: str) -> List[Dict[str, str]]:
    """
    Extrae secuencias de palabras del JSON de Verint y las transforma
    en eventos estructurados con clasificación Interno (Asesor) / Externo (Cliente).
    """
    result_obj = res_data.get("GetInteractionTranscriptionResult") or {}
    data_trans = result_obj.get("Data") or {}
    sequences = data_trans.get("WordsSequences") or []

    events = []
    for seq in sequences:
        if not isinstance(seq, dict):
            continue
        speaker_raw = seq.get("SpeakerName", "")
        is_agent = speaker_raw == "Agent"
        tipo = "Interno" if is_agent else "Externo"
        speaker_label = f"Asesor ({agent_name})" if is_agent else "Cliente"

        start_ms = seq.get("StartTime", 0)
        total_sec = int(start_ms) // 1000
        mins = total_sec // 60
        secs = total_sec % 60
        ts_str = f"{mins:02d}:{secs:02d}"

        words = " ".join([w.get("WordText", "") for w in seq.get("Words", []) if isinstance(w, dict) and w.get("WordText")]).strip()
        if words:
            events.append({
                "tiempo": ts_str,
                "tipo": tipo,
                "emisor": speaker_label,
                "mensaje": words
            })
    return events


def create_interaction_docx(
    call_id: str,
    metadata: Dict[str, Any],
    events: List[Dict[str, str]],
    output_path: Path
) -> Path:
    """
    Construye y guarda el documento Word (.docx) con metadatos y tabla
    compatible con auditores y con el extractor de auditorías wsp_docx_extractor.py.
    """
    doc = docx.Document()

    # Título principal
    title = doc.add_heading(f"Transcripción Verint - Interacción {call_id}", level=1)
    title.runs[0].font.name = "Calibri"

    # 1. Tabla de Metadatos
    doc.add_heading("1. Metadatos de la Interacción", level=2)
    meta_table = doc.add_table(rows=0, cols=2)
    meta_table.style = "Table Grid"

    fields = [
        ("ID Llamada (SWITCH_CALL_ID)", str(call_id)),
        ("Asesor / Colaborador", str(metadata.get("Agent", "No identificado"))),
        ("Fecha y Hora de Grabación", str(metadata.get("LocalStartTime", "No disponible"))),
        ("SID de Contacto", str(metadata.get("SID") or metadata.get("Sid") or metadata.get("DocumentId", "N/A"))),
        ("Total de Intervenciones", str(len(events)))
    ]

    for label, val in fields:
        row_cells = meta_table.add_row().cells
        row_cells[0].text = label
        row_cells[1].text = val
        row_cells[0].paragraphs[0].runs[0].font.bold = True

    doc.add_paragraph()

    # 2. Tabla Tabulada de Diálogo (Compatible con formato Genesys/WSP de auditoría)
    doc.add_heading("2. Detalle de Conversación (Tabulado)", level=2)
    dialog_table = doc.add_table(rows=1, cols=4)
    dialog_table.style = "Table Grid"

    hdr_cells = dialog_table.rows[0].cells
    hdr_cells[0].text = "Tiempo"
    hdr_cells[1].text = "Tipo"
    hdr_cells[2].text = "Emisor"
    hdr_cells[3].text = "Mensaje"
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True

    for ev in events:
        row_cells = dialog_table.add_row().cells
        row_cells[0].text = ev["tiempo"]
        row_cells[1].text = ev["tipo"]
        row_cells[2].text = ev["emisor"]
        row_cells[3].text = ev["mensaje"]

    doc.add_paragraph()

    # 3. Transcripción Continua
    doc.add_heading("3. Diálogo Continuo", level=2)
    for ev in events:
        p = doc.add_paragraph()
        run_ts = p.add_run(f"[{ev['tiempo']}] ")
        run_ts.font.bold = True
        run_ts.font.color.rgb = RGBColor(90, 90, 90)

        run_spk = p.add_run(f"{ev['emisor']}: ")
        run_spk.font.bold = True

        p.add_run(ev["mensaje"])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def export_single_call_to_docx(
    call_id: str,
    client: VerintAPIClient,
    output_dir: Path
) -> Optional[Path]:
    """
    Consulta Verint vía request HTTP para un call_id y genera su .docx correspondiente.
    """
    clean_id = str(call_id).strip()
    logger.info(f"🔎 Consultando transcripción para ID: {clean_id}...")

    res_data = client.get_interaction_transcription_api(clean_id)
    if not res_data:
        logger.warning(f"⚠️ No se encontró transcripción en Verint para ID: {clean_id}")
        return None

    meta = res_data.get("_contact_metadata") or {}
    agent_name = meta.get("Agent", "Desconocido")
    agent_clean = "".join(c for c in agent_name if c.isalnum() or c in (" ", "_", "-")).strip().replace(" ", "_")

    events = extract_call_dialogue_events(res_data, agent_name)
    if not events:
        logger.warning(f"⚠️ Verint retornó registro pero sin secuencias de diálogo para ID: {clean_id}")
        return None

    filename = f"Transcripcion_Verint_{clean_id}_{agent_clean}.docx"
    output_path = output_dir / filename
    saved_path = create_interaction_docx(clean_id, meta, events, output_path)

    logger.info(f"💾 Guardado .docx exitosamente: {saved_path.name} ({len(events)} intervenciones)")
    return saved_path


def load_ids_from_file(file_path: Path) -> List[str]:
    """
    Lee IDs desde archivo de texto (.txt) o columna de Excel (.xlsx).
    """
    if not file_path.exists():
        raise FileNotFoundError(f"No existe el archivo de entrada: {file_path}")

    ids = []
    if file_path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            for val in row:
                if val:
                    val_str = str(val).strip()
                    if val_str.isdigit() and len(val_str) >= 4:
                        ids.append(val_str)
                        break
    else:
        with open(file_path, "r", encoding="utf-8") as f:
            for line in f:
                item = line.strip()
                if item and not item.startswith("#"):
                    ids.append(item)
    return list(dict.fromkeys(ids))


def main():
    load_dotenv()

    parser = argparse.ArgumentParser(description="Exportador de Transcripciones Verint WFO a DOCX por ID de llamada.")
    parser.add_argument("--id", type=str, help="ID único de interacción a consultar.")
    parser.add_argument("--ids", nargs="+", help="Lista de IDs de interacción separados por espacio.")
    parser.add_argument("--file", type=str, help="Ruta a archivo .txt o .xlsx con lista de IDs.")
    parser.add_argument("--output-dir", type=str, default=str(BASE_DIR / "data" / "input" / "auditorias_wsp"), help="Carpeta destino para los archivos .docx.")

    args = parser.parse_args()

    call_ids = []
    if args.id:
        call_ids.append(args.id.strip())
    if args.ids:
        call_ids.extend([cid.strip() for cid in args.ids if cid.strip()])
    if args.file:
        call_ids.extend(load_ids_from_file(Path(args.file)))

    call_ids = list(dict.fromkeys(call_ids))

    if not call_ids:
        logger.error("Debe proporcionar al menos un ID mediante --id, --ids o --file.")
        logger.info("Ejemplo de uso:")
        logger.info("  .\\.venv\\Scripts\\python modules\\verint\\tools\\export_verint_docx_by_ids.py --id 12345678")
        sys.exit(1)

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    user = os.getenv("VERINT_USER")
    pwd = os.getenv("VERINT_PASS")
    if not user or not pwd:
        logger.error("Credenciales VERINT_USER o VERINT_PASS no configuradas en archivo .env.")
        sys.exit(1)

    logger.info("=" * 70)
    logger.info("🎧 EXPORTADOR VERINT A DOCX (CONSULTA HTTP POR ID)")
    logger.info(f"Total IDs a procesar : {len(call_ids)}")
    logger.info(f"Carpeta Destino      : {output_dir}")
    logger.info("=" * 70)

    client = VerintAPIClient(username=user, password=pwd)
    if not client.login():
        logger.error("No se pudo iniciar sesión en Verint WFO. Verifique credenciales o VPN.")
        sys.exit(1)

    exitosos = 0
    fallidos = 0

    try:
        for idx, cid in enumerate(call_ids, 1):
            logger.info(f"[{idx}/{len(call_ids)}] Procesando ID: {cid}")
            try:
                out = export_single_call_to_docx(cid, client, output_dir)
                if out:
                    exitosos += 1
                else:
                    fallidos += 1
            except Exception as e:
                logger.error(f"Error extrayendo ID {cid}: {e}")
                fallidos += 1
    finally:
        client.close()

    logger.info("=" * 70)
    logger.info(f"Proceso finalizado. Exitosos: {exitosos} | Fallidos: {fallidos}")
    logger.info("=" * 70)


if __name__ == "__main__":
    main()
